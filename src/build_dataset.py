"""Flatten 'Training for Tm.xlsx' into a tidy long-format dataset.

One row per (measurement, component). Measurements are identified by meas_id.
Aria's plate-reader DMPC/DPPC series is appended as new measurements.
"""
import os
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
import re
import openpyxl
import pandas as pd

XLSX = os.path.join(ROOT, "data", "raw", "Training for Tm.xlsx")
OUT_LONG = os.path.join(ROOT, "data") + "/measurements_long.csv"
OUT_SMILES = os.path.join(ROOT, "data") + "/lipid_smiles.csv"

PCT = re.compile(r"^(?P<name>.*?)[\s]*(?P<pct>\d+(?:\.\d+)?)\s*%\s*$")
CHAINS = re.compile(r"\((\d+:\d+[^)]*)\)")


def parse_component(cell):
    """'DMPC (14:0) 50%' -> ('DMPC (14:0)', 50.0)"""
    s = str(cell).strip()
    m = PCT.match(s)
    if not m:
        return s, None
    return m.group("name").strip(), float(m.group("pct"))


def normalize(name):
    """Collapse 'DPPC (16:0)' and 'DPPC' to one label so systems group correctly."""
    n = name.strip()
    # Drop a trailing parenthetical that is purely a chain annotation
    n = re.sub(r"\s*\((\d+:\d+[^)]*)\)\s*$", "", n).strip()
    return n


def load_reference_smiles():
    """lipid label -> SMILES from the 'SMILES code reference' sheet."""
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["SMILES code reference"]
    ref = {}
    for r in ws.iter_rows(min_row=2):
        name, smi = r[0].value, r[1].value
        if name in (None, "") or smi in (None, ""):
            continue
        ref[normalize(str(name))] = str(smi).strip()
    return ref


def load_literature():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Aria data"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    recs = []
    for i, r in enumerate(rows[1:], start=2):
        r = r + [None] * (18 - len(r))
        if all(v in (None, "") for v in r[:15]):
            continue
        tm = r[2]
        comps = []
        for ci, si in ((3, 4), (5, 6), (7, 8), (9, 10), (11, 12), (13, 14)):
            if r[ci] in (None, ""):
                continue
            name, pct = parse_component(r[ci])
            comps.append((normalize(name), pct, r[si]))
        if not comps:
            continue
        for name, pct, smiles in comps:
            recs.append(dict(
                meas_id=f"lit_{i}",
                source_row=i,
                lipid=name,
                pct=pct,
                smiles=smiles,
                tm_c=tm,
                method=(str(r[15]).strip() if r[15] not in (None, "") else "unspecified"),
                dataset="literature",
                paper=r[0],
                doi=r[1],
                notes=r[16],
            ))
    return pd.DataFrame(recs)


# --- Aria plate-reader series (DMPC/DPPC), peak of 1st derivative of Laurdan GP fit ---
# Labels as printed on the Prism table. The two repeated "80/20" and "90/10" entries
# are the reversed series (i.e. 20/80 and 10/90 DMPC/DPPC) and are flagged unreliable,
# per Ira: only 100/0 .. 60/40 are considered reliable right now.
ARIA = [
    # (dmpc_pct, dppc_pct, peak_x, reliable, printed_label)
    (100, 0, 32.48, True,  "100/0"),
    (90, 10, 35.28, True,  "90/10"),
    (80, 20, 37.14, True,  "80/20"),
    (70, 30, 38.39, True,  "70/30"),
    (60, 40, 39.81, True,  "60/40"),
    (50, 50, 39.95, False, "50/50"),
    (40, 60, 38.99, False, "40/60"),
    (30, 70, 41.21, False, "30/70"),
    (20, 80, 37.28, False, "80/20 (reversed series)"),
    (10, 90, 38.03, False, "90/10 (reversed series)"),
    (0, 100, 37.58, False, "0/100"),
]

DMPC_SMILES = "CCCCCCCCCCCCCC(=O)OC[C@H](COP(=O)([O-])OCC[N+](C)(C)C)OC(=O)CCCCCCCCCCCCC"
DPPC_SMILES = "CCCCCCCCCCCCCCCC(=O)OC[C@H](COP(=O)([O-])OCC[N+](C)(C)C)OC(=O)CCCCCCCCCCCCCCC"


def load_aria():
    recs = []
    for k, (dm, dp, tm, rel, label) in enumerate(ARIA):
        mid = f"aria_{k:02d}"
        for name, pct, smi in (("DMPC", dm, DMPC_SMILES), ("DPPC", dp, DPPC_SMILES)):
            if pct == 0:
                continue
            recs.append(dict(
                meas_id=mid, source_row=None, lipid=name, pct=float(pct), smiles=smi,
                tm_c=tm, method="Laurdan GP (plate reader)", dataset="aria_2026",
                paper="Aria plate-reader screen, DMPC/DPPC", doi=None,
                notes=f"1st-deriv peak of nonlin fit of GP; label {label}; reliable={rel}",
                reliable=rel,
            ))
    return pd.DataFrame(recs)


def main():
    lit = load_literature()
    lit["reliable"] = True
    aria = load_aria()
    long = pd.concat([lit, aria], ignore_index=True)

    # --- backfill SMILES ---------------------------------------------------
    # Many later literature rows leave the SMILES column blank even though the
    # same lipid is spelled out earlier in the sheet or in the reference tab.
    ref = load_reference_smiles()
    inline = (long.dropna(subset=["smiles"])
                  .drop_duplicates("lipid").set_index("lipid")["smiles"].to_dict())
    lookup = {**ref, **inline}
    n_before = long["smiles"].isna().sum()
    long["smiles"] = long["smiles"].fillna(long["lipid"].map(lookup))
    print(f"SMILES backfilled: {n_before - long['smiles'].isna().sum()} "
          f"rows ({long['smiles'].isna().sum()} still missing)")
    unmapped = sorted(long.loc[long.smiles.isna(), "lipid"].unique())
    if unmapped:
        print("  lipids with NO structure anywhere:", unmapped)
    long["smiles_missing"] = long["smiles"].isna()
    long["meas_smiles_complete"] = ~long.groupby("meas_id")["smiles_missing"].transform("any")

    # sanity: composition sums
    sums = long.groupby("meas_id")["pct"].sum()
    bad = sums[(sums < 95) | (sums > 105)]
    long["pct_sum"] = long["meas_id"].map(sums)
    long["pct_sum_ok"] = ~long["meas_id"].isin(bad.index)

    # mole fraction, renormalised so each measurement sums to exactly 1
    long["mole_frac"] = long["pct"] / long["meas_id"].map(sums)

    # Counterion, parsed out of the Marsh notation in the Notes column. It is
    # already recorded there; it was simply never given to the model. Divalent
    # cations bridge adjacent headgroups and shift Tm enormously: (16:0)4CL
    # melts at 39.5 C as the Na2 salt and 88.3 C as the Ca salt.
    import re as _re
    _ION = _re.compile(r"Marsh notation:.*\.\(?(H|Li|Na|K|NH4|Rb|Cs|Mg|Ca|Ba|Sr)\)?\d*[a-z]?$")
    _VALENCY = {"H": 1, "Li": 1, "Na": 1, "K": 1, "NH4": 1, "Rb": 1, "Cs": 1,
                "Mg": 2, "Ca": 2, "Ba": 2, "Sr": 2}
    ion = long["notes"].astype(str).str.extract(_ION)[0]
    long["counterion"] = ion.fillna("none")
    long["ion_valency"] = long["counterion"].map(_VALENCY).fillna(0).astype(int)
    long["ion_divalent"] = (long["ion_valency"] == 2).astype(int)

    long["n_components"] = long.groupby("meas_id")["lipid"].transform("size")
    long["has_tm"] = long["tm_c"].notna()

    long.to_csv(OUT_LONG, index=False)

    smi = (long.dropna(subset=["smiles"])[["lipid", "smiles"]]
           .drop_duplicates().sort_values("lipid"))
    smi.to_csv(OUT_SMILES, index=False)

    print(f"measurements: {long.meas_id.nunique()}  component-rows: {len(long)}")
    print(f"with Tm: {long[long.has_tm].meas_id.nunique()}")
    print(f"unique lipids: {long.lipid.nunique()}  with SMILES: {smi.lipid.nunique()}")
    print("\ncomponents per measurement:")
    print(long.groupby('meas_id')['n_components'].first().value_counts().sort_index())
    print(f"\nmeasurements usable for a structure-based model "
          f"(Tm present + all SMILES resolved): "
          f"{long[long.has_tm & long.meas_smiles_complete].meas_id.nunique()}")
    print(f"\ncomposition sums outside 95-105%: {len(bad)} measurements")
    if len(bad):
        print(bad.head(15))
    print("\nmethods:")
    print(long.groupby('meas_id').first()['method'].value_counts().head(20))


if __name__ == "__main__":
    main()
